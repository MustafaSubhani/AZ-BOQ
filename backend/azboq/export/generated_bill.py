"""Generated bill: our own lines from the checked items, laid out with the client's template styles.

Row styles are copied from template rows of the same kind (header, section, line,
overhead, total). openpyxl does not auto-size wrapped rows, so each height is
estimated from the description length, the column-B width and the font size.
"""
import math
from copy import copy
from pathlib import Path

import openpyxl

from .. import config

TEMPLATE_SECTION_ROW = 15
TEMPLATE_LINE_ROW = 16
TEMPLATE_TOTAL_ROWS = (2331, 2332, 2333, 2334, 2335)
FIRST = 10


def row_height(text: str, col_width_chars: float, font_pt: float) -> float:
    chars_per_line = max(col_width_chars * 1.15, 10)
    lines = max(1, sum(math.ceil(max(len(part), 1) / chars_per_line) for part in str(text).split("\n")))
    return round(max(15.4, lines * font_pt * 1.32), 2)


def write_generated_bill(bill: dict, report: dict, out: Path) -> dict:
    wb = openpyxl.load_workbook(config.BILL_PATH)
    ws = wb["BOQ"]
    values = openpyxl.load_workbook(config.BILL_PATH, data_only=True)["BOQ"]
    header = [[values.cell(r, c).value for c in range(1, 13)] for r in range(1, FIRST)]
    total_styles = {r: [copy(ws.cell(r, c)._style) for c in range(1, 13)] for r in TEMPLATE_TOTAL_ROWS}
    total_labels = {r: values.cell(r, 2).value for r in TEMPLATE_TOTAL_ROWS}
    sig = [copy(ws.cell(2337, c)._style) for c in range(1, 13)]
    section_styles = [copy(ws.cell(TEMPLATE_SECTION_ROW, c)._style) for c in range(1, 13)]
    line_styles = [copy(ws.cell(TEMPLATE_LINE_ROW, c)._style) for c in range(1, 13)]
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= FIRST:
            ws.unmerge_cells(str(mr))
    ws.conditional_formatting = type(ws.conditional_formatting)()
    ws.data_validations.dataValidation = []
    ws.delete_rows(FIRST, ws.max_row - FIRST + 1)
    for r, row in enumerate(header, start=1):
        for c, v in enumerate(row, start=1):
            if not isinstance(ws.cell(r, c), openpyxl.cell.cell.MergedCell):
                ws.cell(r, c).value = v
    col_b = ws.column_dimensions["B"].width or 60.8
    by_id = {l["id"]: l for l in bill["lines"]}
    r = FIRST
    written = 0
    for s in bill["sections"]:
        rows = [by_id[x] for x in s["lines"] if report["lines"][x]["status"] in ("matched", "different", "printed", "derived")]
        if not rows:
            continue
        for col in range(1, 13):
            ws.cell(r, col)._style = copy(section_styles[col - 1])
        ws.cell(r, 1).value = s["code"]
        ws.cell(r, 2).value = s["title"]
        # The template's external-link cells (C and L here, J and L on lines) print their cached value 0.
        ws.cell(r, 3).value = 0
        ws.cell(r, 12).value = 0
        ws.row_dimensions[r].height = row_height(s["title"], col_b, 11)
        r += 1
        for i, l in enumerate(rows, start=1):
            res = report["lines"][l["id"]]
            for col in range(1, 13):
                ws.cell(r, col)._style = copy(line_styles[col - 1])
            qty = res.get("our_qty") if res.get("our_qty") is not None else l["qty"]
            ws.cell(r, 1).value = i
            ws.cell(r, 2).value = l["description"]
            ws.cell(r, 3).value = l["unit_raw"]
            ws.cell(r, 4).value = qty
            ws.cell(r, 5).value = l["rate"]
            ws.cell(r, 6).value = f"=+E{r}*D{r}"
            ws.cell(r, 7).value = f"=I{r}+K{r}"
            ws.cell(r, 8).value = f"=J{r}+L{r}"
            ws.cell(r, 9).value = f"=J{r}/F{r}"
            ws.cell(r, 10).value = 0
            ws.cell(r, 11).value = f"=L{r}/F{r}"
            ws.cell(r, 12).value = 0
            ws.row_dimensions[r].height = row_height(l["description"], col_b, 11)
            written += 1
            r += 1
    last_line = r - 1
    labels = list(TEMPLATE_TOTAL_ROWS)
    formulas = {
        2331: f"=SUM(F{FIRST}:F{last_line})*1.8%", 2332: f"=SUM(F{FIRST}:F{last_line + 1})*2%",
        2333: f"=SUM(F{FIRST}:F{last_line + 2})", 2334: f"=F{last_line + 3}*18%", 2335: f"=F{last_line + 3}+F{last_line + 4}",
    }
    for k, src in enumerate(labels):
        for col in range(1, 13):
            ws.cell(r, col)._style = copy(total_styles[src][col - 1])
        ws.cell(r, 2).value = total_labels[src]
        ws.cell(r, 6).value = formulas[src]
        ws.row_dimensions[r].height = 18
        r += 1
    r += 1
    for col in range(1, 13):
        ws.cell(r, col)._style = copy(sig[col - 1])
    ws.cell(r, 2).value = "Employer / Sifarişçi"
    ws.cell(r, 6).value = "Contractor / Podratçı"
    ws["C5"].value = f"=F{last_line + 5}"
    ws.print_area = f"A1:L{r}"
    wb._external_links = []
    for name in list(wb.defined_names.keys()):
        del wb.defined_names[name]
    wb.calculation.fullCalcOnLoad = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"path": str(out), "lines": written, "last_row": r}

