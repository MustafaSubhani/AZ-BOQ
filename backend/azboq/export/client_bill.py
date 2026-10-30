"""Filled client bill: the original template, values only where we have a measured quantity.

openpyxl strips cached values from formulas that point at other workbooks, and the
workbook behind [6] is not in the package. So every external formula is replaced
by its cached value first, the external links and broken defined names are
removed, and only then is the file saved. Excel recalculates on open.
"""
import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment

from .. import config, runlog

EXTERNAL = re.compile(r"\[\d+\]")


def write_client_bill(out_path: Path, quantities: dict[int, float] | None = None, notes: dict[int, str] | None = None,
                      evidence_rows: list[dict] | None = None, highlight: bool = False) -> dict:
    quantities = quantities or {}
    notes = notes or {}
    src = config.BILL_PATH
    wf = openpyxl.load_workbook(src, data_only=False)
    wv = openpyxl.load_workbook(src, data_only=True)
    ws, vs = wf["BOQ"], wv["BOQ"]
    replaced = 0
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and EXTERNAL.search(c.value):
                c.value = vs[c.coordinate].value
                replaced += 1
    # Rows without a stored height show at the sheet default in the client's Excel.
    # Replacing formulas with values makes Excel auto-fit them, which adds 11 pages
    # to the print, so they are pinned at the default height.
    default_h = ws.sheet_format.defaultRowHeight
    pinned = 0
    for r in range(1, ws.max_row + 1):
        dim = ws.row_dimensions[r]
        if dim.height is None:
            dim.height = default_h
            pinned += 1
    changed = []
    for r, qty in quantities.items():
        cell = ws.cell(r, 4)
        if cell.value != qty:
            changed.append(r)
        cell.value = qty
        if highlight:
            from openpyxl.styles import PatternFill
            cell.fill = PatternFill("solid", fgColor="FFF1EA")
    for r, text in notes.items():
        ws.cell(r, 4).comment = Comment(text, "AZ-BOQ check")
    removed_names = []
    for name in list(wf.defined_names.keys()):
        dn = wf.defined_names[name]
        if dn.attr_text and EXTERNAL.search(dn.attr_text):
            del wf.defined_names[name]
            removed_names.append(name)
    for sheet in wf.worksheets:
        for name in list(sheet.defined_names.keys()):
            if EXTERNAL.search(sheet.defined_names[name].attr_text or ""):
                del sheet.defined_names[name]
                removed_names.append(f"{sheet.title}!{name}")
    wf._external_links = []
    wf.calculation.fullCalcOnLoad = True
    if evidence_rows is not None:
        ev = wf.create_sheet("Evidence")
        headers = ["Line", "Row", "Status", "Our quantity", "Bill quantity", "Unit", "Method", "Source", "Reason", "Review"]
        ev.append(headers)
        for e in evidence_rows:
            ev.append([e.get(h.lower().replace(" ", "_")) for h in headers])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wf.save(out_path)
    runlog.event("export", "info", f"client bill written: {replaced} external formulas replaced by cached values, "
                 f"{len(removed_names)} broken names removed, {len(changed)} quantities changed")
    return {"path": str(out_path), "external_replaced": replaced, "names_removed": removed_names,
            "changed_rows": changed, "rows_pinned_to_default_height": pinned}


def run_bounded(work, timeout_s: int, kill) -> dict:
    """Run work(box) and kill its process if it exceeds timeout_s. box['pid'] is the process to stop."""
    import threading
    box: dict = {}

    def target():
        try:
            box["result"] = work(box)
        except Exception as exc:
            box["error"] = exc

    thread = threading.Thread(target=target, daemon=True)
    thread.start()
    thread.join(timeout_s)
    if thread.is_alive():
        kill(box.get("pid"))
        raise TimeoutError(f"Excel did not finish within {timeout_s}s")
    if "error" in box:
        raise box["error"]
    return box.get("result") or {}


def _excel_pid(excel) -> int | None:
    import ctypes
    raw = getattr(excel, "Hwnd", None)
    if raw is None:
        return None
    hwnd = int(raw)
    if hwnd == 0:
        return None
    pid = ctypes.c_ulong()
    ctypes.windll.user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
    value = int(pid.value)
    if value == 0:
        return None
    return value


def excel_finish(xlsx: Path, pdf: Path | None = None, timeout_s: int = 300) -> dict:
    """Recalculate and save with real Excel, then print to PDF. Windows only.

    If Excel hangs, only the Excel process this call started is stopped.
    """
    def kill(pid):
        if not pid:
            return
        import subprocess
        subprocess.run(["taskkill", "/PID", str(pid), "/F"], check=False, capture_output=True)

    def work(box):
        import pythoncom
        import win32com.client as win32
        pythoncom.CoInitialize()
        excel = win32.DispatchEx("Excel.Application")
        box["pid"] = _excel_pid(excel)
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AskToUpdateLinks = False
        result = {}
        try:
            wb = excel.Workbooks.Open(str(xlsx.resolve()), UpdateLinks=0)
            excel.CalculateFull()
            wb.Save()
            if pdf is not None:
                pdf.parent.mkdir(parents=True, exist_ok=True)
                wb.Worksheets("BOQ").ExportAsFixedFormat(0, str(pdf.resolve()))
                result["pdf"] = str(pdf)
            result["links"] = list(wb.LinkSources(1) or [])
            wb.Close(SaveChanges=False)
        finally:
            excel.Quit()
            pythoncom.CoUninitialize()
        return result

    return run_bounded(work, timeout_s, kill)
