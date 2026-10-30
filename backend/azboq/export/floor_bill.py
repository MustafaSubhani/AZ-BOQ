"""Complete bill, same template, rows gathered by the storey the text names.

Every client line is kept, with the client's quantity, so the total is the contract
total. A line moves under Zirzəmi, 1-ci mərtəbə, 2-ci mərtəbə or Dam only when its
own text or its section title says so. The rest stay under "Floor not stated".
This file is separate from BOQ-Hadrut-checked.xlsx, which keeps the client's row order.
"""
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment

from .. import config
from ..floors import UNSTATED, label
from .generated_bill import FIRST, TEMPLATE_LINE_ROW, TEMPLATE_SECTION_ROW, TEMPLATE_TOTAL_ROWS, row_height


def write_floor_bill(bill: dict, report: dict, floors: dict, out: Path) -> dict:
    wb = openpyxl.load_workbook(config.BILL_PATH)
    ws = wb["BOQ"]
    values = openpyxl.load_workbook(config.BILL_PATH, data_only=True)["BOQ"]
    header = [[values.cell(r, c).value for c in range(1, 13)] for r in range(1, FIRST)]
    total_styles = {r: [copy(ws.cell(r, c)._style) for c in range(1, 13)] for r in TEMPLATE_TOTAL_ROWS}
    total_labels = {r: values.cell(r, 2).value for r in TEMPLATE_TOTAL_ROWS}
    sig = [copy(ws.cell(2337, c)._style) for c in range(1, 13)]
    section_styles = [copy(ws.cell(TEMPLATE_SECTION_ROW, c)._style) for c in range(1, 13)]
    line_styles = [copy(ws.cell(TEMPLATE_LINE_ROW, c)._style) for c in range(1, 13)]
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= FIRST:
            ws.unmerge_cells(str(mr))
    ws.conditional_formatting = type(ws.conditional_formatting)()
    ws.data_validations.dataValidation = []
    ws.delete_rows(FIRST, ws.max_row - FIRST + 1)
    for r, row in enumerate(header, start=1):
        for c, v in enumerate(row, start=1):
            if not isinstance(ws.cell(r, c), openpyxl.cell.cell.MergedCell):
                ws.cell(r, c).value = v
    col_b = ws.column_dimensions["B"].width or 60.8  # allow-default: column width, not a quantity
    by_id = {l["id"]: l for l in bill["lines"]}
    sections = {s["code"]: s for s in bill["sections"]}
    grouped: dict[str, list] = {fid: [] for fid in floors["order"]}
    for line in bill["lines"]:
        fid = floors["lines"].get(line["id"], {}).get("floor", UNSTATED)
        grouped.setdefault(fid, []).append(line)

    r = FIRST
    written = 0
    noted = 0
    for fid in floors["order"]:
        rows = grouped.get(fid) or []
        if not rows:
            continue
        az, en = label(fid)
        _section_row(ws, r, section_styles, f"{az}  /  {en}", col_b)
        r += 1
        current = None
        for line in rows:
            if line["section"] != current:
                current = line["section"]
                sec = sections[line["section"]]
                _section_row(ws, r, section_styles, f"{sec['code']}  {sec['title']}", col_b, code=sec["code"])
                r += 1
            res = report["lines"][line["id"]]
            for col in range(1, 13):
                ws.cell(r, col)._style = copy(line_styles[col - 1])
            ws.cell(r, 1).value = line["item"]
            ws.cell(r, 2).value = line["description"]
            ws.cell(r, 3).value = line["unit_raw"]
            ws.cell(r, 4).value = line["qty"]
            ws.cell(r, 5).value = line["rate"]
            ws.cell(r, 6).value = f"=+E{r}*D{r}"
            ws.cell(r, 7).value = f"=I{r}+K{r}"
            ws.cell(r, 8).value = f"=J{r}+L{r}"
            ws.cell(r, 9).value = f"=J{r}/F{r}"
            ws.cell(r, 10).value = 0
            ws.cell(r, 11).value = f"=L{r}/F{r}"
            ws.cell(r, 12).value = 0
            ws.row_dimensions[r].height = row_height(line["description"], col_b, 11)
            ours = res.get("our_qty")
            if ours is not None and res.get("status") in ("matched", "different"):
                ws.cell(r, 4).comment = Comment(
                    f"Drawing quantity {ours:g} {line['unit']} ({res['status']}). {res.get('reason') or ''}",
                    "AZ-BOQ",
                )
                noted += 1
            written += 1
            r += 1
    last_line = r - 1
    formulas = {
        2331: f"=SUM(F{FIRST}:F{last_line})*1.8%",
        2332: f"=SUM(F{FIRST}:F{last_line + 1})*2%",
        # The client's subtotal is the lines, the 1.8% and the 2%, minus 0.02. Row 10 is a heading, so its amount cell is empty.
        2333: f"=SUM(F{FIRST}:F{last_line + 2})-0.02",
        2334: f"=F{last_line + 3}*18%",
        2335: f"=F{last_line + 3}+F{last_line + 4}",
    }
    for src in TEMPLATE_TOTAL_ROWS:
        for col in range(1, 13):
            ws.cell(r, col)._style = copy(total_styles[src][col - 1])
        ws.cell(r, 2).value = total_labels[src]
        ws.cell(r, 6).value = formulas[src]
        ws.row_dimensions[r].height = 18
        r += 1
    r += 1
    for col in range(1, 13):
        ws.cell(r, col)._style = copy(sig[col - 1])
    ws.cell(r, 2).value = "Employer / Sifarişçi"
    ws.cell(r, 6).value = "Contractor / Podratçı"
    ws["C5"].value = f"=F{last_line + 5}"
    ws.print_area = f"A1:L{r}"
    wb._external_links = []
    for name in list(wb.defined_names.keys()):
        del wb.defined_names[name]
    wb.calculation.fullCalcOnLoad = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"path": str(out), "lines": written, "drawing_notes": noted, "last_row": r,
            "by_id_present": written == len(by_id)}


def _section_row(ws, r, styles, title, col_b, code=None):
    for col in range(1, 13):
        ws.cell(r, col)._style = copy(styles[col - 1])
    ws.cell(r, 1).value = code
    ws.cell(r, 2).value = title
    ws.cell(r, 3).value = 0
    ws.cell(r, 12).value = 0
    ws.row_dimensions[r].height = row_height(title, col_b, 11)
