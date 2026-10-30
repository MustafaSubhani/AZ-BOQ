"""Compare a written bill with the client's original: every cell's style, the sheet layout and the print setup."""
import openpyxl
import pymupdf

LAST_ROW, LAST_COL = 2339, 12


def _cell_sig(c) -> tuple:
    f, a, b, p = c.font, c.alignment, c.border, c.protection
    fill = c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None
    return (
        f.name, f.sz, f.b, f.i, f.u, str(f.color.rgb) if f.color is not None and f.color.type == "rgb" else None,
        fill, c.number_format, a.horizontal, a.vertical, a.wrap_text, a.indent,
        b.left.style, b.right.style, b.top.style, b.bottom.style, p.locked,
    )


def compare_styles(original, written) -> dict:
    wo = openpyxl.load_workbook(original).active
    ww = openpyxl.load_workbook(written)["BOQ"]
    diffs = []
    for r in range(1, LAST_ROW + 1):
        for col in range(1, LAST_COL + 1):
            a, b = wo.cell(r, col), ww.cell(r, col)
            if _cell_sig(a) != _cell_sig(b):
                diffs.append({"cell": a.coordinate, "original": _cell_sig(a), "written": _cell_sig(b)})
    layout = []
    default_h = wo.sheet_format.defaultRowHeight
    for r in range(1, LAST_ROW + 1):
        ho = wo.row_dimensions[r].height
        ho = default_h if ho is None else ho
        hw = ww.row_dimensions[r].height
        hw = default_h if hw is None else hw
        # Excel stores heights rounded to its pixel grid (31.25 is saved as 31.2).
        if abs(ho - hw) > 0.1:
            layout.append(f"row {r} height {ho} vs {hw}")
    for k, d in wo.column_dimensions.items():
        if round(d.width or 0, 2) != round(ww.column_dimensions[k].width or 0, 2):  # allow-default: missing column width, not a quantity
            layout.append(f"column {k} width {d.width} vs {ww.column_dimensions[k].width}")
    if sorted(map(str, wo.merged_cells.ranges)) != sorted(map(str, ww.merged_cells.ranges)):
        layout.append("merged ranges differ")
    for attr in ("print_area", "print_title_rows", "freeze_panes"):
        if getattr(wo, attr) != getattr(ww, attr):
            layout.append(f"{attr}: {getattr(wo, attr)} vs {getattr(ww, attr)}")
    for attr in ("orientation", "paperSize", "fitToWidth", "fitToHeight"):
        if getattr(wo.page_setup, attr) != getattr(ww.page_setup, attr):
            layout.append(f"page_setup.{attr}: {getattr(wo.page_setup, attr)} vs {getattr(ww.page_setup, attr)}")
    for attr in ("left", "right", "top", "bottom"):
        if getattr(wo.page_margins, attr) != getattr(ww.page_margins, attr):
            layout.append(f"margin {attr}")
    if wo.sheet_properties.pageSetUpPr.fitToPage != ww.sheet_properties.pageSetUpPr.fitToPage:
        layout.append("fitToPage differs")
    if len(wo.conditional_formatting) != len(ww.conditional_formatting):
        layout.append("conditional formats differ")
    if len(wo.data_validations.dataValidation) != len(ww.data_validations.dataValidation):
        layout.append("data validations differ")
    return {"cell_style_diffs": len(diffs), "examples": diffs[:10], "layout_diffs": layout}


def compare_values(original, written, changed_rows: set[int]) -> dict:
    vo = openpyxl.load_workbook(original, data_only=True).active
    vw = openpyxl.load_workbook(written, data_only=True)["BOQ"]
    diffs = []
    for r in range(1, LAST_ROW + 1):
        if r in changed_rows:
            continue
        for col in range(1, LAST_COL + 1):
            a, b = vo.cell(r, col).value, vw.cell(r, col).value
            if isinstance(a, float) and isinstance(b, (int, float)):
                if abs(a - b) > 1e-6 * max(1.0, abs(a)):  # allow-default: value comparison epsilon, not a quantity
                    diffs.append((vo.cell(r, col).coordinate, a, b))
            elif a != b and not (a in (None, "") and b in (None, "")):
                diffs.append((vo.cell(r, col).coordinate, a, b))
    return {"value_diffs": len(diffs), "examples": diffs[:10]}


def compare_print(original_pdf, written_pdf, pages=(0, 29), dpi=60) -> dict:
    a, b = pymupdf.open(original_pdf), pymupdf.open(written_pdf)
    out = {"original_pages": a.page_count, "written_pages": b.page_count,
           "written_size": [round(b[0].rect.width), round(b[0].rect.height)], "pixel_diff": {}}
    for p in pages:
        if p >= min(a.page_count, b.page_count):
            continue
        pa = a[p].get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
        pb = b[p].get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
        if (pa.width, pa.height) != (pb.width, pb.height):
            out["pixel_diff"][p] = "page sizes differ"
            continue
        sa, sb = pa.samples, pb.samples
        changed = sum(1 for x, y in zip(sa, sb) if abs(x - y) > 48)
        out["pixel_diff"][p] = round(changed / len(sa), 5)
    return out
